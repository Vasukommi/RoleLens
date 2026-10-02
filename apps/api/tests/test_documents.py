from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter

from rolelens.documents import DocumentError, make_passages, parse_resume


def test_text_upload_preserves_source_and_sanitizes_filename():
    text = "Synthetic engineer\n\nBuilt React dashboards using TypeScript."
    resume = parse_resume("../../resume.txt", text.encode())
    assert resume.filename == "resume.txt"
    assert resume.text == text
    assert all(passage.text in resume.text for passage in resume.passages)


def test_docx_includes_table_text():
    document = Document()
    document.add_paragraph("Synthetic engineer with application development experience.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "React"
    table.cell(0, 1).text = "Built customer dashboards"
    buffer = BytesIO()
    document.save(buffer)
    result = parse_resume("resume.docx", buffer.getvalue())
    assert "React | Built customer dashboards" in result.text


@pytest.mark.parametrize(
    "filename,content",
    [
        ("resume.txt", b""),
        ("resume.exe", b"Synthetic engineer with application development experience"),
        ("resume.pdf", b"not a PDF"),
        ("resume.docx", b"not a DOCX"),
        ("resume.txt", b"\xff\xfeinvalid text"),
    ],
)
def test_invalid_documents_are_rejected(filename, content):
    with pytest.raises(DocumentError):
        parse_resume(filename, content)


def test_scanned_or_blank_pdf_is_not_silently_assessed():
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=600)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="OCR"):
        parse_resume("scan.pdf", buffer.getvalue())


def test_compressed_docx_bomb_is_rejected_before_parsing():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * (16 * 1024 * 1024))
    with pytest.raises(DocumentError, match="archive is too large"):
        parse_resume("resume.docx", buffer.getvalue())


def test_large_text_is_rejected_without_truncating_evidence():
    with pytest.raises(DocumentError, match="24,000"):
        parse_resume("resume.txt", b"x" * 24_001)


def test_many_paragraphs_remain_within_choice_limit():
    text = "\n\n".join(f"Experience statement {index}" for index in range(500))
    passages = make_passages(text)
    assert len(passages) <= 200
    assert all(passage.text in text for passage in passages)
    assert "Experience statement 499" in passages[-1].text
