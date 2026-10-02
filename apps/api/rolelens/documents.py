import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from rolelens.ocr import OcrError, extract_pages
from rolelens.schemas import MAX_RESUME_CHARS, ParsedResume, Passage

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_DOCX_UNPACKED_BYTES = 15 * 1024 * 1024


class DocumentError(ValueError):
    pass


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < 30:
        raise DocumentError(
            "No usable resume text found. Scanned PDFs need OCR; try a text PDF or DOCX."
        )
    if len(text) > MAX_RESUME_CHARS:
        raise DocumentError("Resume text exceeds 24,000 characters. Upload a shorter document.")
    return text


def make_passages(text: str) -> list[Passage]:
    # Verbatim substrings: model-selected IDs always resolve to original source text.
    chunks = []
    cursor = 0
    while cursor < len(text):
        end = min(cursor + 600, len(text))
        paragraph_end = text.find("\n\n", cursor, end)
        if paragraph_end >= 0:
            end = paragraph_end
        elif end < len(text):
            newline = text.rfind("\n", cursor + 100, end)
            space = text.rfind(" ", cursor + 100, end)
            boundary = max(newline, space)
            if boundary >= 0:
                end = boundary
        chunk = text[cursor:end].strip()
        if chunk:
            chunks.append(Passage(id=f"p{len(chunks) + 1}", text=chunk))
        cursor = end
        while cursor < len(text) and text[cursor].isspace():
            cursor += 1
    if len(chunks) > 200:
        # Keep the evidence choice below Jev's 255-option limit without dropping source text.
        return [
            Passage(id=f"p{i + 1}", text=text[start : start + 600])
            for i, start in enumerate(range(0, len(text), 600))
        ]
    return chunks


def parse_resume(filename: str, content: bytes) -> ParsedResume:
    if not content:
        raise DocumentError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise DocumentError("Files must be 5 MB or smaller.")
    safe_name = Path(filename.replace("\\", "/")).name[:180]
    extension = Path(safe_name).suffix.lower()
    if extension not in {".pdf", ".docx", ".txt"}:
        raise DocumentError("Upload a PDF, DOCX, or UTF-8 TXT resume.")
    extraction_method = "native"
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise DocumentError("Password-protected PDFs are not supported.")
            if len(reader.pages) > 20:
                raise DocumentError("Resume PDFs must have 20 pages or fewer.")
            texts = [page.extract_text() or "" for page in reader.pages]
            missing = [index for index, text in enumerate(texts) if len(text.strip()) < 30]
            if missing:
                recovered = extract_pages(content, missing)
                texts = [recovered.get(index, text) for index, text in enumerate(texts)]
                extraction_method = "ocr" if len(missing) == len(texts) else "mixed"
            text = "\n\n".join(texts)
        elif extension == ".docx":
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 256
                    or sum(e.file_size for e in entries) > MAX_DOCX_UNPACKED_BYTES
                ):
                    raise DocumentError("The DOCX archive is too large to process.")
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(paragraphs)
        else:
            text = content.decode("utf-8-sig")
    except DocumentError:
        raise
    except OcrError as error:
        raise DocumentError(str(error)) from error
    except (BadZipFile, UnicodeDecodeError) as error:
        raise DocumentError(
            "The document is invalid or uses an unsupported text encoding."
        ) from error
    except Exception as error:
        raise DocumentError(
            "Could not read this document. Try exporting a new PDF or DOCX."
        ) from error
    text = normalize_text(text)
    return ParsedResume(
        filename=safe_name,
        text=text,
        passages=make_passages(text),
        extraction_method=extraction_method,
    )
